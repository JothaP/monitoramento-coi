"""Consulta Operacional — Ambiente de Teste (Teresina).

Versão V3. Substitui SOMENTE pages/Ambiente_de_Teste.py.
- Apenas administrador autenticado e autorizado a usar "teste".
- Base de matrículas XLS/XLSX ou SQLite indexado, somente leitura.
- Índice SQLite local criado SOMENTE na primeira pesquisa caso necessário.
- Nenhum histórico de consultas, nem escrita no Excel ou em dados de produção.
- Regiões de influência extraídas exclusivamente dos polígonos KMZ/KML.
"""
from __future__ import annotations

from collections import defaultdict
from html import escape as html_escape
import hashlib
import math
import os
from pathlib import Path
import re
import sqlite3
import sys
import tempfile
import threading
import unicodedata
import uuid
import zipfile
import xml.etree.ElementTree as ET

import folium
from folium.plugins import MarkerCluster
import pandas as pd
import streamlit as st
from streamlit_folium import st_folium

# Compatibilidade com importação de auth.py no Streamlit Cloud.
RAIZ = Path(__file__).resolve().parent.parent
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))
from auth import verificar_autenticacao, tem_acesso_modulo

st.set_page_config(
    page_title="Consulta Operacional | Ambiente de Teste - COI",
    page_icon="🔍",
    layout="wide",
    initial_sidebar_state="collapsed",
)
if not verificar_autenticacao():
    st.error("Sessão não autenticada. Acesse o Hub Central para entrar.")
    st.stop()
if (str(st.session_state.get("perfil", "")).strip().lower() != "admin"
        or not tem_acesso_modulo("teste")):
    st.error("A Consulta Operacional em desenvolvimento é exclusiva do administrador.")
    st.stop()

PASTA_DADOS = RAIZ / "dados"
ARQUIVO_POCOS = PASTA_DADOS / "pocos.csv"
ARQUIVO_LOGGERS = PASTA_DADOS / "loggers.csv"
PASTA_BAIRROS = RAIZ / "geodados" / "municipios"
PASTA_INFLUENCIAS = RAIZ / "geodados" / "areas_influencia"

COLUNAS_MATRICULAS = (
    "CATEGORIA", "NUM_LIGACAO", "BAIRRO",
    "COD_LATITUDE", "COD_LONGITUDE", "END_LIGACAO",
)
COLUNAS_POCOS = (
    "ID_POCO", "IDENTIFICACAO_ATIVO", "NOME_POCO", "MUNICIPIO", "LATITUDE", "LONGITUDE",
)
COLUNAS_LOGGERS = (
    "ID_LOGGER", "IDENTIFICACAO_ATIVO", "ENDERECO", "MUNICIPIO", "LATITUDE", "LONGITUDE",
)

# Processo de indexação é único no processo Streamlit e não é feito no carregamento da página.
_BLOQUEIO_INDICE = threading.Lock()


def _normalizar_texto(valor: object) -> str:
    texto = unicodedata.normalize("NFKD", str(valor if valor is not None else "").strip().upper())
    return "".join(c for c in texto if not unicodedata.combining(c))


def _normalizar_cabecalho(valor: object) -> str:
    return re.sub(r"[^A-Z0-9]", "", _normalizar_texto(valor))


def _matricula(valor: object) -> str:
    """Não remove zeros à esquerda que estejam armazenados como texto."""
    valor = str(valor if valor is not None else "").strip()
    if re.fullmatch(r"\d+\.0+", valor):
        return valor.split(".", 1)[0]
    return valor


def _coordenada(valor: object, tipo: str) -> float | None:
    if valor is None or pd.isna(valor):
        return None
    bruto = str(valor).strip().replace(" ", "")
    if not bruto:
        return None
    if "," in bruto and "." in bruto:
        bruto = (bruto.replace(".", "").replace(",", ".")
                 if bruto.rfind(",") > bruto.rfind(".") else bruto.replace(",", ""))
    else:
        bruto = bruto.replace(",", ".")
    try:
        numero = float(bruto)
    except (ValueError, TypeError, OverflowError):
        return None
    if not math.isfinite(numero) or numero == 0:
        return None
    if tipo == "lat" and not -90 <= numero <= 90:
        return None
    if tipo == "lon" and not -180 <= numero <= 180:
        return None
    return numero


def _assinatura(caminho: Path) -> tuple[str, int, int]:
    info = caminho.stat()
    return str(caminho), info.st_size, info.st_mtime_ns


def _arquivo_matriculas() -> Path | None:
    if not PASTA_DADOS.is_dir():
        return None
    arquivos = {p.name.lower(): p for p in PASTA_DADOS.iterdir() if p.is_file()}
    return arquivos.get("matriculas.xls") or arquivos.get("matriculas.xlsx")


def _arquivo_sqlite_publicado() -> Path | None:
    """Arquivo indexado opcional previamente produzido pelo conversor."""
    if not PASTA_DADOS.is_dir():
        return None
    arquivos = {p.name.lower(): p for p in PASTA_DADOS.iterdir() if p.is_file()}
    return arquivos.get("matriculas.sqlite") or arquivos.get("matriculas.db")


@st.cache_data(show_spinner=False, max_entries=4)
def _resumo_hash(assinatura: tuple[str, int, int]) -> str:
    """Hash de leitura sequencial, sem interpretar planilha nem carregar tudo na RAM."""
    h = hashlib.sha256()
    with open(assinatura[0], "rb") as arquivo:
        for bloco in iter(lambda: arquivo.read(2 * 1024 * 1024), b""):
            h.update(bloco)
    return h.hexdigest()


def _conexao_leitura(db: Path):
    """Abre o SQLite SEM permitir gravação pelo mecanismo de consulta."""
    return sqlite3.connect(db.resolve().as_uri() + "?mode=ro", uri=True)


def _metadados_indice(db: Path) -> dict[str, str] | None:
    try:
        with _conexao_leitura(db) as con:
            meta = dict(con.execute("SELECT chave, valor FROM metadados"))
            # Evita considerar um arquivo incompleto como base válida.
            con.execute("SELECT NUM_LIGACAO FROM matriculas LIMIT 1").fetchone()
            return meta if meta.get("versao") == "3" else None
    except (sqlite3.Error, OSError, ValueError):
        return None


def _linhas_excel(caminho: Path):
    """Percorre TODAS as abas com cabeçalho esperado sem montar um DataFrame gigante.

    Detecta o formato real pelo cabeçalho do arquivo, inclusive arquivo XLSX
    indevidamente nomeado .xls. Um .xls antigo suporta no máximo 65.536 linhas por aba.
    """
    with caminho.open("rb") as f:
        assinatura = f.read(8)
    if assinatura.startswith(b"PK"):
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise RuntimeError("Adicione openpyxl ao requirements.txt.") from e
        # Arquivo binário, não caminho: aceita conteúdo XLSX mesmo se nomeado .xls.
        entrada = caminho.open("rb")
        try:
            workbook = load_workbook(entrada, read_only=True, data_only=True)
            for aba in workbook.worksheets:
                linhas = aba.iter_rows(values_only=True)
                cabecalho = next(linhas, None)
                if cabecalho is None:
                    continue
                campos = {_normalizar_cabecalho(v): i for i, v in enumerate(cabecalho)}
                if any(_normalizar_cabecalho(c) not in campos for c in COLUNAS_MATRICULAS):
                    continue
                indices = [campos[_normalizar_cabecalho(c)] for c in COLUNAS_MATRICULAS]
                yield aba.title, (tuple((linha[i] if i < len(linha) and linha[i] is not None else "")
                                      for i in indices) for linha in linhas)
        finally:
            if "workbook" in locals():
                workbook.close()
            entrada.close()
    elif assinatura.startswith(b"\xd0\xcf\x11\xe0"):
        try:
            import xlrd
        except ImportError as e:
            raise RuntimeError("Adicione xlrd>=2.0.1 ao requirements.txt para ler .xls.") from e
        workbook = xlrd.open_workbook(str(caminho), on_demand=True)
        try:
            for aba in workbook.sheets():
                if aba.nrows < 1:
                    continue
                campos = {_normalizar_cabecalho(v): i for i, v in enumerate(aba.row_values(0))}
                if any(_normalizar_cabecalho(c) not in campos for c in COLUNAS_MATRICULAS):
                    continue
                indices = [campos[_normalizar_cabecalho(c)] for c in COLUNAS_MATRICULAS]
                def obter(worksheet=aba, cols=indices):
                    for r in range(1, worksheet.nrows):
                        yield tuple(worksheet.cell_value(r, i) if i < worksheet.ncols else "" for i in cols)
                yield aba.name, obter()
        finally:
            workbook.release_resources()
    else:
        raise ValueError("Formato Excel desconhecido. Envie uma planilha .xls ou .xlsx válida.")


def _texto_celula(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _construir_indice_excel(excel: Path, destino: Path, digest: str) -> int:
    """Índice local descartável. Nunca grava sobre o Excel nem sobre base existente."""
    destino.parent.mkdir(parents=True, exist_ok=True)
    temporario = destino.with_name(destino.name + "." + uuid.uuid4().hex + ".tmp")
    total = 0
    abas_validas = 0
    try:
        os.close(os.open(temporario, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600))
        con = sqlite3.connect(temporario)
        try:
            con.execute("PRAGMA journal_mode=OFF")
            con.execute("PRAGMA synchronous=OFF")
            con.execute("CREATE TABLE matriculas ("
                        "CATEGORIA TEXT, NUM_LIGACAO TEXT NOT NULL, BAIRRO TEXT, "
                        "COD_LATITUDE TEXT, COD_LONGITUDE TEXT, END_LIGACAO TEXT)")
            con.execute("CREATE TABLE metadados (chave TEXT PRIMARY KEY, valor TEXT)")
            lote = []
            for nome_aba, linhas in _linhas_excel(excel):
                abas_validas += 1
                for linha in linhas:
                    valores = [_texto_celula(c) for c in linha]
                    valores[1] = _matricula(valores[1])
                    if not valores[1]:
                        continue
                    lote.append(valores)
                    if len(lote) >= 5000:
                        con.executemany("INSERT INTO matriculas VALUES (?,?,?,?,?,?)", lote)
                        total += len(lote)
                        lote.clear()
                if lote:
                    con.executemany("INSERT INTO matriculas VALUES (?,?,?,?,?,?)", lote)
                    total += len(lote)
                    lote.clear()
            if not abas_validas:
                raise ValueError("Nenhuma aba tem as seis colunas esperadas do cadastro de matrículas.")
            con.execute("CREATE INDEX idx_num_ligacao ON matriculas(NUM_LIGACAO)")
            con.executemany("INSERT INTO metadados(chave,valor) VALUES (?,?)", [
                ("versao", "3"), ("fonte", excel.name), ("sha256", digest),
                ("quantidade", str(total)), ("abas_validas", str(abas_validas)),
            ])
            con.commit()
        finally:
            con.close()
        os.replace(temporario, destino)
        os.chmod(destino, 0o600)
        return total
    finally:
        if temporario.exists():
            temporario.unlink()


def _indice_matriculas() -> tuple[Path, dict[str, str]]:
    """Executada somente após o clique em Pesquisar.

    Usa base SQLite publicada e verificada; senão índice local em /tmp,
    gerado uma única vez por versão da planilha durante a vida do ambiente.
    """
    fonte = _arquivo_matriculas()
    publicado = _arquivo_sqlite_publicado()
    if fonte is None and publicado is None:
        raise FileNotFoundError("Adicione dados/matriculas.xls, .xlsx ou matriculas.sqlite.")

    if fonte is None:
        meta = _metadados_indice(publicado)
        if meta is None:
            raise ValueError(f"O índice {publicado.name} é inválido. Gere-o novamente.")
        return publicado, meta

    assinatura = _assinatura(fonte)
    digest = _resumo_hash(assinatura)
    if publicado is not None:
        meta = _metadados_indice(publicado)
        if meta and meta.get("sha256") == digest:
            return publicado, meta

    codigo = hashlib.sha256((str(fonte.resolve()) + "|" + str(assinatura[1]) + "|" + digest).encode()).hexdigest()[:20]
    destino = Path(tempfile.gettempdir()) / "coi_consulta_operacional" / f"matriculas_{codigo}.sqlite"
    meta = _metadados_indice(destino)
    if meta and meta.get("sha256") == digest:
        return destino, meta
    with _BLOQUEIO_INDICE:
        meta = _metadados_indice(destino)
        if not meta or meta.get("sha256") != digest:
            _construir_indice_excel(fonte, destino, digest)
        meta = _metadados_indice(destino)
    if meta is None:
        raise RuntimeError("Não foi possível gerar o índice de matrículas.")
    return destino, meta


def _buscar_matriculas(codigos: list[str]) -> tuple[pd.DataFrame, dict[str, str]]:
    caminho, meta = _indice_matriculas()
    if not codigos:
        return pd.DataFrame(columns=COLUNAS_MATRICULAS), meta
    with _conexao_leitura(caminho) as con:
        sql = ("SELECT CATEGORIA, NUM_LIGACAO, BAIRRO, COD_LATITUDE, COD_LONGITUDE, END_LIGACAO "
               f"FROM matriculas WHERE NUM_LIGACAO IN ({','.join('?' for _ in codigos)})")
        base = pd.read_sql_query(sql, con, params=codigos, dtype=str).fillna("")
    # Converter apenas coordenadas dos registros solicitados.
    base["_lat"] = base["COD_LATITUDE"].map(lambda v: _coordenada(v, "lat"))
    base["_lon"] = base["COD_LONGITUDE"].map(lambda v: _coordenada(v, "lon"))
    ordem = {codigo: i for i, codigo in enumerate(codigos)}
    if not base.empty:
        base["_ord"] = base["NUM_LIGACAO"].map(ordem)
        base = base.sort_values("_ord", kind="stable").drop(columns="_ord")
    return base, meta


def _validar_colunas(df: pd.DataFrame, esperadas: tuple[str, ...], origem: str) -> pd.DataFrame:
    mapa = {_normalizar_cabecalho(c): c for c in df.columns}
    faltantes = [c for c in esperadas if _normalizar_cabecalho(c) not in mapa]
    if faltantes:
        raise ValueError(f"{origem}: colunas não encontradas: {', '.join(faltantes)}.")
    selecionadas = [mapa[_normalizar_cabecalho(c)] for c in esperadas]
    resultado = df.loc[:, selecionadas].copy()
    resultado.columns = list(esperadas)
    return resultado.fillna("")


@st.cache_data(show_spinner=False, max_entries=6)
def _ler_ativos_csv(assinatura: tuple[str, int, int], tipo: str) -> pd.DataFrame:
    caminho = assinatura[0]
    colunas = COLUNAS_POCOS if tipo == "Poço / Ativo" else COLUNAS_LOGGERS
    try:
        df = pd.read_csv(caminho, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError(f"{Path(caminho).name}: salve o CSV como UTF-8.") from exc
    dados = _validar_colunas(df, colunas, Path(caminho).name)
    dados["_lat"] = dados["LATITUDE"].map(lambda v: _coordenada(v, "lat"))
    dados["_lon"] = dados["LONGITUDE"].map(lambda v: _coordenada(v, "lon"))
    dados = dados.loc[dados["_lat"].notna() & dados["_lon"].notna()].copy()
    dados["_tipo"] = tipo
    dados["_nome"] = dados["NOME_POCO"] if tipo == "Poço / Ativo" else dados["ENDERECO"]
    return dados


def _distance_m(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    r = 6371000.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp, dl = math.radians(lat2 - lat1), math.radians(lon2 - lon1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(min(1, math.sqrt(a)))


def _grade_espacial(df: pd.DataFrame, celula: float = 0.01):
    grade: dict[tuple[int, int], list[tuple[int, float, float]]] = defaultdict(list)
    registros = df.to_dict("records")
    for indice, registro in enumerate(registros):
        lat, lon = float(registro["_lat"]), float(registro["_lon"])
        grade[math.floor(lat / celula), math.floor(lon / celula)].append((indice, lat, lon))
    return grade, registros


def _buscar_proximos(grade, registros, lat: float, lon: float, raio_m: float, celula: float = 0.01):
    delta_lat = raio_m / 110000.0 + 1e-8
    cosseno = max(0.000001, math.cos(math.radians(min(89.999, abs(lat) + delta_lat))))
    delta_lon = raio_m / (110000.0 * cosseno) + 1e-8
    for ix in range(math.floor((lat - delta_lat) / celula), math.floor((lat + delta_lat) / celula) + 1):
        for iy in range(math.floor((lon - delta_lon) / celula), math.floor((lon + delta_lon) / celula) + 1):
            for indice, alat, alon in grade.get((ix, iy), ()):
                distancia = _distance_m(lat, lon, alat, alon)
                if distancia <= raio_m:
                    yield indice, registros[indice], distancia


def _consultar_proximidade(clientes: pd.DataFrame, bases: dict[str, pd.DataFrame], raio: int):
    resultados = []
    ids_mapa: set[tuple[str, str, str]] = set()
    ativos_mapa = []
    for tipo, base in bases.items():
        if base.empty:
            continue
        grade, registros = _grade_espacial(base)
        for _, cliente in clientes.iterrows():
            lat, lon = float(cliente["_lat"]), float(cliente["_lon"])
            for _, ativo, distancia in _buscar_proximos(grade, registros, lat, lon, raio):
                identificacao = str(ativo.get("IDENTIFICACAO_ATIVO", "")).strip()
                nome = str(ativo.get("_nome", "")).strip()
                resultados.append({
                    "Matrícula": cliente["NUM_LIGACAO"],
                    "Tipo": tipo,
                    "Identificação": identificacao,
                    "Nome / endereço": nome,
                    "Distância (m)": round(distancia, 1),
                })
                chave = (tipo, str(ativo.get("ID_POCO") or ativo.get("ID_LOGGER") or identificacao),
                         f"{ativo['_lat']:.6f},{ativo['_lon']:.6f}")
                if chave not in ids_mapa:
                    ids_mapa.add(chave)
                    ativos_mapa.append(ativo)
    resultado_df = pd.DataFrame(resultados,
        columns=["Matrícula", "Tipo", "Identificação", "Nome / endereço", "Distância (m)"])
    if not resultado_df.empty:
        resultado_df = resultado_df.sort_values(["Matrícula", "Distância (m)", "Tipo"], kind="stable")
    return resultado_df, ativos_mapa


def _interpretar_pesquisa(texto: str) -> tuple[list[str], list[str]]:
    valores = re.split(r"[,;\s]+", texto.strip())
    vistos = set()
    matriculas = []
    invalidos = []
    for valor in valores:
        if not valor:
            continue
        numero = _matricula(valor)
        if not numero.isdecimal():
            invalidos.append(valor)
        elif numero not in vistos:
            vistos.add(numero)
            matriculas.append(numero)
    return matriculas, invalidos


def _popup_cliente(linha: pd.Series) -> str:
    """Exatamente os três campos pedidos; escape contra HTML nas planilhas e no KML."""
    campos = (
        ("Matrícula", linha.get("NUM_LIGACAO", "")),
        ("Endereço", linha.get("END_LIGACAO", "")),
        ("Área de Influência", linha.get("_influencia", "Área de influência não identificada")),
    )
    return "<div style='min-width:245px'>" + "".join(
        f"<div><b>{html_escape(str(k))}:</b> {html_escape(str(v))}</div>" for k, v in campos
    ) + "</div>"


def _popup_ativo(ativo: dict) -> str:
    return "<div style='min-width:190px'>" + "".join(
        f"<div><b>{html_escape(k)}:</b> {html_escape(str(v))}</div>"
        for k, v in (("Tipo", ativo.get("_tipo", "")),
                     ("Identificação", ativo.get("IDENTIFICACAO_ATIVO", "")),
                     ("Nome / endereço", ativo.get("_nome", "")))
    ) + "</div>"


def _caminho_bairros_teresina() -> Path | None:
    if not PASTA_BAIRROS.is_dir():
        return None
    candidatos = [p for p in PASTA_BAIRROS.iterdir()
                  if p.is_file() and p.suffix.lower() in {".gpkg", ".kml", ".kmz"}
                  and "TERESINA" in _normalizar_texto(p.stem)
                  and not any(t in _normalizar_texto(p.stem) for t in ("INFLUENCIA", "ABRANGENCIA"))]
    prioridade = {".gpkg": 0, ".kmz": 1, ".kml": 2}
    if candidatos:
        return sorted(candidatos, key=lambda p: (prioridade[p.suffix.lower()], p.name))[0]
    legado = RAIZ / "TERESINA.kmz"
    return legado if legado.is_file() else None


@st.cache_data(show_spinner=False, max_entries=2)
def _bairros_geograficos(assinatura: tuple[str, int, int]) -> list[dict]:
    """Código anterior preservado para as feições de bairro, independente da influência."""
    from pyproj import CRS, Transformer
    from shapely import wkb
    from shapely.geometry import Polygon
    from shapely.ops import transform

    caminho = Path(assinatura[0])
    bairros = []
    if caminho.suffix.lower() == ".gpkg":
        with sqlite3.connect(f"file:{caminho}?mode=ro", uri=True) as banco:
            tabelas = banco.execute("SELECT table_name, srs_id FROM gpkg_contents WHERE data_type='features'").fetchall()
            for tabela, srs_id in tabelas:
                campo_geo = banco.execute(
                    "SELECT column_name, srs_id FROM gpkg_geometry_columns WHERE table_name=?", (tabela,)
                ).fetchone()
                if not campo_geo:
                    continue
                nomes = [c[1] for c in banco.execute('PRAGMA table_info("' + tabela.replace('"', '""') + '")')]
                campos = {_normalizar_cabecalho(c): c for c in nomes}
                campo_nome = next((campos[c] for c in ("NMBAIRRO", "NOMEBAIRRO", "BAIRRO", "NOME") if c in campos), None)
                campo_municipio = next((campos[c] for c in ("NMMUN", "MUNICIPIO", "NOMEMUNICIPIO", "CIDADE") if c in campos), None)
                if not campo_nome:
                    continue
                nome_geom, srs_geom = campo_geo
                linha_crs = banco.execute(
                    "SELECT organization, organization_coordsys_id, definition FROM gpkg_spatial_ref_sys WHERE srs_id=?",
                    (srs_geom or srs_id,),
                ).fetchone()
                if not linha_crs:
                    continue
                org, cod, wkt_crs = linha_crs
                try:
                    origem = CRS.from_user_input(f"{org}:{cod}") if org and cod else CRS.from_wkt(wkt_crs)
                except Exception:
                    origem = CRS.from_wkt(wkt_crs)
                conversor = None if origem.equals(CRS.from_epsg(4326)) else Transformer.from_crs(origem, 4326, always_xy=True)
                citar = lambda s: '"' + s.replace('"', '""') + '"'
                cols = [nome_geom, campo_nome] + ([campo_municipio] if campo_municipio else [])
                consulta = f"SELECT {', '.join(citar(c) for c in cols)} FROM {citar(tabela)}"
                for registro in banco.execute(consulta):
                    blob, nome, *mun = registro
                    if not blob or not nome or (mun and mun[0] and _normalizar_texto(mun[0]) != "TERESINA"):
                        continue
                    blob = bytes(blob)
                    if blob[:2] != b"GP" or len(blob) < 8:
                        continue
                    envelope = (blob[3] >> 1) & 7
                    tamanhos = {0: 0, 1: 32, 2: 48, 3: 48, 4: 64}
                    if envelope not in tamanhos:
                        continue
                    geo = wkb.loads(blob[8 + tamanhos[envelope]:])
                    if conversor:
                        geo = transform(conversor.transform, geo)
                    for area in ([geo] if geo.geom_type == "Polygon" else list(geo.geoms) if geo.geom_type == "MultiPolygon" else []):
                        if not area.is_empty:
                            simples = area.simplify(0.00007, preserve_topology=True)
                            visual = simples if simples.geom_type == "Polygon" else area
                            bairros.append({"nome": str(nome), "coords": [(y, x) for x, y in visual.exterior.coords],
                                            "limites": area.bounds, "geometria": area})
    else:
        if caminho.suffix.lower() == ".kmz":
            with zipfile.ZipFile(caminho) as zipado:
                kmls = [n for n in zipado.namelist() if n.lower().endswith(".kml")]
                if not kmls:
                    return []
                xml = zipado.read(kmls[0])
        else:
            xml = caminho.read_bytes()
        root = ET.fromstring(xml)
        ns = {"k": "http://www.opengis.net/kml/2.2"}
        for lugar in root.findall(".//k:Placemark", ns):
            el_nome = lugar.find("k:name", ns)
            nome = el_nome.text.strip() if el_nome is not None and el_nome.text else ""
            if not nome:
                continue
            for poligono in lugar.findall(".//k:Polygon", ns):
                ring = poligono.find("k:outerBoundaryIs/k:LinearRing/k:coordinates", ns)
                if ring is None or not ring.text:
                    continue
                pontos = []
                for termo in ring.text.split():
                    try:
                        longitude, latitude = (float(x) for x in termo.split(",")[:2])
                        pontos.append((longitude, latitude))
                    except (ValueError, TypeError):
                        continue
                if len(pontos) >= 3:
                    area = Polygon(pontos)
                    if not area.is_empty:
                        simples = area.simplify(0.00007, preserve_topology=True)
                        visual = simples if simples.geom_type == "Polygon" else area
                        bairros.append({"nome": nome, "coords": [(y, x) for x, y in visual.exterior.coords],
                                        "limites": area.bounds, "geometria": area})
    return bairros


def _caminho_influencias() -> Path | None:
    """Prefere diretório dedicado; evita confundir polígonos de bairros e abastecimento."""
    candidatos = []
    if PASTA_INFLUENCIAS.is_dir():
        candidatos.extend(p for p in PASTA_INFLUENCIAS.iterdir()
                          if p.is_file() and p.suffix.lower() in {".kmz", ".kml"})
    if not candidatos and PASTA_BAIRROS.is_dir():
        candidatos.extend(p for p in PASTA_BAIRROS.iterdir()
                          if p.is_file() and p.suffix.lower() in {".kmz", ".kml"}
                          and ("INFLUENCIA" in _normalizar_texto(p.stem)
                               or "ABRANGENCIA" in _normalizar_texto(p.stem)))
    return sorted(candidatos, key=lambda p: p.name)[0] if candidatos else None


def _coordenadas_kml(texto: str | None) -> list[tuple[float, float]]:
    pontos = []
    for item in (texto or "").split():
        partes = item.split(",")
        if len(partes) < 2:
            continue
        try:
            lon, lat = float(partes[0]), float(partes[1])
        except ValueError:
            continue
        if math.isfinite(lon) and math.isfinite(lat) and -180 <= lon <= 180 and -90 <= lat <= 90:
            pontos.append((lon, lat))
    return pontos


@st.cache_data(show_spinner=False, max_entries=3)
def _ler_influencias(assinatura: tuple[str, int, int]) -> list[dict]:
    """Carrega todos os polígonos, inclusive MultiGeometry e buracos internos.

    Não insere no mapa os pontos do KML, apenas polígonos das áreas.
    Geometria original é preservada para cruzamento exato; simplificação só visual.
    """
    from shapely.geometry import Polygon
    caminho = Path(assinatura[0])
    if caminho.suffix.lower() == ".kmz":
        with zipfile.ZipFile(caminho) as compactado:
            arquivos = [n for n in compactado.namelist() if n.lower().endswith(".kml")]
            if not arquivos:
                raise ValueError("O KMZ não contém arquivos KML.")
            fontes = [compactado.read(n) for n in arquivos]
    else:
        fontes = [caminho.read_bytes()]
    ns = {"k": "http://www.opengis.net/kml/2.2"}
    saida = []
    for xml in fontes:
        raiz = ET.fromstring(xml)
        for placemark in raiz.findall(".//k:Placemark", ns):
            elemento_nome = placemark.find("k:name", ns)
            nome = (elemento_nome.text or "").strip() if elemento_nome is not None else ""
            if not nome:
                continue
            for elemento in placemark.findall(".//k:Polygon", ns):
                externa = elemento.find("k:outerBoundaryIs/k:LinearRing/k:coordinates", ns)
                vertices = _coordenadas_kml(externa.text if externa is not None else None)
                if len(vertices) < 3:
                    continue
                furos = []
                for anel in elemento.findall("k:innerBoundaryIs/k:LinearRing/k:coordinates", ns):
                    coords = _coordenadas_kml(anel.text)
                    if len(coords) >= 3:
                        furos.append(coords)
                poligono = Polygon(vertices, holes=furos)
                if not poligono.is_valid:
                    poligono = poligono.buffer(0)
                if poligono.is_empty:
                    continue
                partes = ([poligono] if poligono.geom_type == "Polygon"
                          else list(poligono.geoms) if poligono.geom_type == "MultiPolygon" else [])
                for parte in partes:
                    visual = parte.simplify(0.00007, preserve_topology=True)
                    if visual.is_empty or visual.geom_type != "Polygon":
                        visual = parte
                    saida.append({
                        "nome": nome,
                        "geometria": parte,
                        "limites": parte.bounds,
                        "contorno": [[lat, lon] for lon, lat in visual.exterior.coords],
                        "buracos": [[[lat, lon] for lon, lat in anel.coords] for anel in visual.interiors],
                    })
    return saida


def _identificar_influencias(lat: float, lon: float, areas: list[dict]) -> list[str]:
    from shapely.geometry import Point
    ponto = Point(lon, lat)
    correspondentes = []
    vistos = set()
    for area in areas:
        minlon, minlat, maxlon, maxlat = area["limites"]
        if not (minlon <= lon <= maxlon and minlat <= lat <= maxlat):
            continue
        if area["geometria"].covers(ponto):
            nome = area["nome"]
            if nome not in vistos:
                vistos.add(nome)
                correspondentes.append(nome)
    return correspondentes


def _mapa_consulta(clientes: pd.DataFrame, ativos: list[dict], raio: int,
                   bairros: list[dict], mostrar_bairros: bool, areas: list[dict]) -> folium.Map:
    latitudes = clientes["_lat"].astype(float).tolist()
    longitudes = clientes["_lon"].astype(float).tolist()
    mapa = folium.Map(location=[sum(latitudes) / len(latitudes), sum(longitudes) / len(longitudes)],
                      zoom_start=14, control_scale=True, tiles="OpenStreetMap")

    if mostrar_bairros and bairros:
        grupo = folium.FeatureGroup(name="Limites de bairros", show=True)
        for bairro in bairros:
            folium.Polygon(
                locations=bairro["coords"], color="#578bba", weight=1.5,
                fill=True, fill_color="#88b8de", fill_opacity=0.06,
                tooltip=f"Bairro: {bairro['nome']}",
            ).add_to(grupo)
        grupo.add_to(mapa)

    # Nova camada separada dos bairros. Transparência para não ocultar marcadores.
    if areas:
        grupo_influencias = folium.FeatureGroup(name="Áreas de influência", show=True)
        for area in areas:
            folium.Polygon(
                locations=([area["contorno"]] + area["buracos"]
                           if area["buracos"] else area["contorno"]),
                color="#8762ab", weight=1.7, fill=True,
                fill_color="#8762ab", fill_opacity=0.11,
                tooltip=f"Área de influência: {area['nome']}",
            ).add_to(grupo_influencias)
        grupo_influencias.add_to(mapa)

    grupo_clientes = folium.FeatureGroup(name="Matrículas", show=True)
    aglomerador = MarkerCluster(name="Matrículas", overlay=False, control=False).add_to(grupo_clientes)
    for _, linha in clientes.iterrows():
        lat, lon = float(linha["_lat"]), float(linha["_lon"])
        folium.Marker(
            [lat, lon], tooltip=f"Matrícula: {linha['NUM_LIGACAO']}",
            popup=folium.Popup(_popup_cliente(linha), max_width=370),
            icon=folium.Icon(color="green", icon="home", prefix="fa"),
        ).add_to(aglomerador)
        folium.Circle([lat, lon], radius=raio, color="#227ab6", fill=False,
                      weight=1, opacity=0.30, tooltip=f"Raio de busca: {raio} m").add_to(grupo_clientes)
    grupo_clientes.add_to(mapa)

    # Mesmos ícones. Mudam SOMENTE as cores dos marcadores.
    for tipo, cor, icone in (("Poço / Ativo", "blue", "tint"), ("Logger", "orange", "signal")):
        grupo = folium.FeatureGroup(name=tipo + " próximo(s)", show=True)
        for ativo in ativos:
            if ativo.get("_tipo") != tipo:
                continue
            folium.Marker(
                [float(ativo["_lat"]), float(ativo["_lon"])],
                tooltip=f"{tipo}: {ativo.get('IDENTIFICACAO_ATIVO', '')}",
                popup=folium.Popup(_popup_ativo(ativo), max_width=360),
                icon=folium.Icon(color=cor, icon=icone, prefix="fa"),
            ).add_to(grupo)
        grupo.add_to(mapa)

    if len(clientes) > 1:
        mapa.fit_bounds([[min(latitudes), min(longitudes)], [max(latitudes), max(longitudes)]], padding=(45, 45))
    folium.LayerControl(collapsed=False).add_to(mapa)
    return mapa


def exibir_consulta_operacional() -> None:
    """Pesquisa e mapa são gerados só depois de o usuário clicar em Pesquisar."""
    if not verificar_autenticacao() or not tem_acesso_modulo("teste") or str(st.session_state.get("perfil", "")).strip().lower() != "admin":
        st.error("A Consulta Operacional em desenvolvimento é exclusiva do administrador.")
        st.stop()

    st.title("🔍 Consulta Operacional")
    st.caption("🧪 Protótipo no Ambiente de Teste • Teresina • Consultas somente leitura • Sem histórico")

    fonte = _arquivo_matriculas()
    indice = _arquivo_sqlite_publicado()
    if fonte is None and indice is None:
        st.error("Base ausente. Adicione dados/matriculas.xls, dados/matriculas.xlsx "
                 "ou dados/matriculas.sqlite a um repositório privado.")
        st.stop()
    if fonte is not None:
        st.caption(f"Base de consulta: dados/{fonte.name} • leitura somente após pesquisar")
    else:
        st.caption(f"Base de consulta: dados/{indice.name} • índice de pesquisa")

    # Descarta apenas o valor de demonstração herdado da versão anterior,
    # sem limpar a consulta atual em futuras atualizações da interface.
    if "consulta_operacional_v3_inicializou" not in st.session_state:
        st.session_state.pop("consulta_operacional_entrada", None)
        st.session_state.pop("consulta_operacional_consulta", None)
        st.session_state["consulta_operacional_v3_inicializou"] = True

    with st.form("consulta_operacional_busca"):
        pesquisa = st.text_area("Matrícula(s)", value="", placeholder="Digite uma ou várias matrículas",
                                help="Informe uma ou várias matrículas, separadas por vírgula, espaço ou quebra de linha.",
                                height=105, key="consulta_operacional_entrada")
        c1, c2 = st.columns([2, 1])
        with c1:
            raio = st.number_input("Raio de busca dos ativos e loggers (metros)",
                                   min_value=50, max_value=5000, value=750, step=50,
                                   key="consulta_operacional_raio")
        with c2:
            exibir_bairros = st.checkbox("Exibir limites dos bairros", value=True,
                                        key="consulta_operacional_bairros")
        executar = st.form_submit_button("🔎 Pesquisar matrículas", type="primary", use_container_width=True)

    if executar:
        codigos, invalidos_pesquisa = _interpretar_pesquisa(pesquisa)
        if invalidos_pesquisa:
            st.error("Valores inválidos na pesquisa: " + ", ".join(invalidos_pesquisa[:10]) + ". Use apenas números de matrícula.")
            st.session_state.pop("consulta_operacional_consulta", None)
        elif not codigos:
            st.warning("Informe ao menos uma matrícula para consultar.")
            st.session_state.pop("consulta_operacional_consulta", None)
        elif len(codigos) > 200:
            st.warning("Para manter o mapa responsivo, pesquise até 200 matrículas por consulta nesta versão de testes.")
            st.session_state.pop("consulta_operacional_consulta", None)
        else:
            # Somente o estado da consulta corrente no navegador; NÃO registra histórico.
            st.session_state["consulta_operacional_consulta"] = (codigos, int(raio), bool(exibir_bairros))

    consulta = st.session_state.get("consulta_operacional_consulta")
    if not consulta:
        st.info("Informe a(s) matrícula(s) e clique em Pesquisar para visualizar o mapa e os ativos próximos.")
        return

    matriculas, raio_consulta, mostrar_bairros = consulta
    try:
        with st.spinner("Consultando cadastro de matrículas... Na primeira pesquisa, a criação do índice pode demorar."):
            encontrados, meta = _buscar_matriculas(matriculas)
    except (OSError, ValueError, RuntimeError, ImportError, sqlite3.Error) as erro:
        st.error(f"Não foi possível consultar a base de matrículas: {erro}")
        return

    st.caption(f"Cadastro indexado: {int(meta.get('quantidade', 0)):,} registros".replace(",", "."))
    presentes = set(encontrados["NUM_LIGACAO"])
    ausentes = [m for m in matriculas if m not in presentes]
    validos = encontrados.loc[encontrados["_lat"].notna() & encontrados["_lon"].notna()].copy()
    invalidos = encontrados.loc[encontrados["_lat"].isna() | encontrados["_lon"].isna()].copy()

    c1, c2, c3 = st.columns(3)
    c1.metric("Matrículas localizadas", len(presentes))
    c2.metric("Com coordenadas válidas", int(validos["NUM_LIGACAO"].nunique()))
    c3.metric("Não encontradas", len(ausentes))

    if ausentes:
        st.warning("Matrícula(s) não encontrada(s): " + ", ".join(ausentes))
    if not invalidos.empty:
        st.warning("Encontradas sem coordenadas válidas: " + ", ".join(sorted(invalidos["NUM_LIGACAO"].unique())))
    duplicadas = encontrados["NUM_LIGACAO"][encontrados["NUM_LIGACAO"].duplicated(keep=False)].unique()
    if len(duplicadas):
        st.warning("Matrículas com registros duplicados na base: " + ", ".join(duplicadas))

    if not encontrados.empty:
        st.subheader("📋 Dados cadastrais das matrículas")
        st.dataframe(encontrados.loc[:, list(COLUNAS_MATRICULAS)], hide_index=True,
                     use_container_width=True)
    if validos.empty:
        st.info("Nenhuma matrícula com coordenadas geográficas válidas para exibir no mapa.")
        return

    # Independente do filtro de limites de bairros: influência sempre é calculada.
    areas = []
    caminho_areas = _caminho_influencias()
    if caminho_areas:
        try:
            areas = _ler_influencias(_assinatura(caminho_areas))
        except Exception as erro:
            st.warning(f"Não foi possível carregar áreas de influência: {erro}")
    else:
        st.warning("KMZ de influência não encontrado: adicione geodados/areas_influencia/Area_de_influencia_Teresina.kmz.")

    if areas:
        validos["_influencia"] = validos.apply(
            lambda linha: "; ".join(_identificar_influencias(float(linha["_lat"]), float(linha["_lon"]), areas))
            or "Área de influência não identificada", axis=1)
    else:
        validos["_influencia"] = "Área de influência indisponível" if caminho_areas is None or not areas else "Área de influência não identificada"

    bases: dict[str, pd.DataFrame] = {}
    for tipo, caminho in (("Poço / Ativo", ARQUIVO_POCOS), ("Logger", ARQUIVO_LOGGERS)):
        if not caminho.is_file():
            st.warning(f"{caminho.relative_to(RAIZ)} não encontrado; {tipo.lower()}(s) não serão exibidos.")
            continue
        try:
            bases[tipo] = _ler_ativos_csv(_assinatura(caminho), tipo)
        except (OSError, ValueError, pd.errors.ParserError) as erro:
            st.warning(f"Não foi possível carregar {caminho.name}: {erro}")

    proximidade, ativos_mapa = _consultar_proximidade(validos, bases, raio_consulta)
    qtd_pocos = sum(a.get("_tipo") == "Poço / Ativo" for a in ativos_mapa)
    qtd_loggers = sum(a.get("_tipo") == "Logger" for a in ativos_mapa)
    c4, c5, c6 = st.columns(3)
    c4.metric("Poços / ativos próximos", qtd_pocos)
    c5.metric("Loggers próximos", qtd_loggers)
    c6.metric("Raio de pesquisa", f"{raio_consulta} m")

    bairros = []
    if mostrar_bairros:
        arquivo_bairros = _caminho_bairros_teresina()
        if arquivo_bairros:
            try:
                bairros = _bairros_geograficos(_assinatura(arquivo_bairros))
            except Exception as erro:
                st.warning(f"Não foi possível carregar os limites dos bairros: {erro}")
        else:
            st.caption("Camada de bairros indisponível: nenhum arquivo geográfico de Teresina foi encontrado.")

    st.subheader("🗺️ Contexto geográfico das matrículas")
    st.caption("Verde: matrículas • Azul: poços / ativos • Laranja: loggers • "
               "Círculos: raio de proximidade. Proximidade não comprova abastecimento pelo ativo.")
    mapa = _mapa_consulta(validos, ativos_mapa, raio_consulta, bairros, mostrar_bairros, areas)
    st_folium(mapa, width=None, height=640,
              returned_objects=[], key="consulta_operacional_mapa")

    st.subheader("📍 Ativos e loggers próximos por matrícula")
    if proximidade.empty:
        st.info(f"Nenhum ativo ou logger encontrado dentro de {raio_consulta} m das matrículas localizadas.")
    else:
        st.dataframe(proximidade, hide_index=True, use_container_width=True)
        st.caption("As distâncias são medidas em linha reta entre coordenadas; "
                   "o mesmo equipamento pode aparecer vinculado a mais de uma matrícula.")

    st.subheader("🌐 Região de influência do abastecimento")
    if areas:
        st.dataframe(validos[["NUM_LIGACAO", "_influencia"]].rename(columns={
            "NUM_LIGACAO": "Matrícula", "_influencia": "Área de Influência"
        }), hide_index=True, use_container_width=True)
        st.caption("A associação resulta da interseção geográfica com os polígonos do KMZ. "
                   "Quando há sobreposições, todas as áreas correspondentes são apresentadas.")
    else:
        st.info("A base de áreas de influência ainda não está disponível para cruzamento geográfico.")
    st.caption("Cadastro de rede: funcionalidade futura, fora do escopo deste protótipo.")


exibir_consulta_operacional()
