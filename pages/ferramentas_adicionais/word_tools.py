"""Word Tools — Formatador de Relatórios ABNT/Corporativo para Streamlit.

Adaptação do pipeline V1.1 fornecido pelo usuário, mantendo seus três
perfis, a classificação de títulos, o modo conservador e a auditoria.
Nenhum arquivo enviado é salvo em disco ou enviado ao Google Colab.
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from collections import Counter
from dataclasses import dataclass

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ESTILOS_PROTEGIDOS = {
    "Heading 1", "Heading 2", "Heading 3", "Caption", "List Paragraph",
    "Quote", "Intense Quote", "TOC Heading", "Bibliography", "Subtitle",
    "List", "List Bullet", "List Number", "Header", "Footer",
}

PERFIS = {
    "ABNT_10719": {
        "page": {"width": 21.0, "height": 29.7, "top": 3.0, "left": 3.0, "bottom": 2.0, "right": 2.0},
        "styles": {
            "Normal": {"font": "Arial", "size": 12, "bold": False, "spacing": 1.5, "indent": 1.25, "align": WD_ALIGN_PARAGRAPH.JUSTIFY},
            "Heading 1": {"font": "Arial", "size": 14, "bold": True, "spacing": 1.5, "indent": 0.0, "align": WD_ALIGN_PARAGRAPH.LEFT, "all_caps": True},
            "Heading 2": {"font": "Arial", "size": 12, "bold": True, "spacing": 1.5, "indent": 0.0, "align": WD_ALIGN_PARAGRAPH.LEFT, "all_caps": False},
            "Heading 3": {"font": "Arial", "size": 12, "bold": False, "spacing": 1.5, "indent": 0.0, "align": WD_ALIGN_PARAGRAPH.LEFT, "all_caps": False},
            "Caption": {"font": "Arial", "size": 10, "bold": False, "spacing": 1.0, "indent": 0.0, "align": WD_ALIGN_PARAGRAPH.CENTER},
        },
    },
    "ABNT_14724": {
        "page": {"width": 21.0, "height": 29.7, "top": 3.0, "left": 3.0, "bottom": 2.0, "right": 2.0},
        "styles": {
            "Normal": {"font": "Arial", "size": 12, "bold": False, "spacing": 1.5, "indent": 1.25, "align": WD_ALIGN_PARAGRAPH.JUSTIFY},
            "Heading 1": {"font": "Arial", "size": 12, "bold": True, "spacing": 1.5, "indent": 0.0, "align": WD_ALIGN_PARAGRAPH.LEFT, "all_caps": True},
            "Heading 2": {"font": "Arial", "size": 12, "bold": False, "spacing": 1.5, "indent": 0.0, "align": WD_ALIGN_PARAGRAPH.LEFT, "all_caps": False},
            "Heading 3": {"font": "Arial", "size": 12, "bold": False, "spacing": 1.5, "indent": 0.0, "align": WD_ALIGN_PARAGRAPH.LEFT, "all_caps": False},
            "Caption": {"font": "Arial", "size": 10, "bold": False, "spacing": 1.0, "indent": 0.0, "align": WD_ALIGN_PARAGRAPH.CENTER},
        },
    },
    "CORPORATIVO": {
        "page": {"width": 21.0, "height": 29.7, "top": 2.5, "left": 2.5, "bottom": 2.5, "right": 2.5},
        "styles": {
            "Normal": {"font": "Calibri", "size": 11, "bold": False, "spacing": 1.15, "indent": 0.0, "align": WD_ALIGN_PARAGRAPH.LEFT},
            "Heading 1": {"font": "Calibri", "size": 16, "bold": True, "spacing": 1.15, "indent": 0.0, "align": WD_ALIGN_PARAGRAPH.LEFT, "all_caps": False},
            "Heading 2": {"font": "Calibri", "size": 14, "bold": True, "spacing": 1.15, "indent": 0.0, "align": WD_ALIGN_PARAGRAPH.LEFT, "all_caps": False},
            "Heading 3": {"font": "Calibri", "size": 12, "bold": True, "spacing": 1.15, "indent": 0.0, "align": WD_ALIGN_PARAGRAPH.LEFT, "all_caps": False},
            "Caption": {"font": "Calibri", "size": 9, "bold": True, "spacing": 1.0, "indent": 0.0, "align": WD_ALIGN_PARAGRAPH.CENTER},
        },
    },
}

ROTULOS_PERFIS = {
    "ABNT_10719": "ABNT 10719 — Relatório técnico",
    "ABNT_14724": "ABNT 14724 — Trabalho acadêmico",
    "CORPORATIVO": "Corporativo — Relatório executivo",
}

MAX_BYTES = 25 * 1024 * 1024


def _limpar_texto(texto: str) -> str:
    return re.sub(r"\s+", "", texto)


def _sha(texto: str) -> str:
    return hashlib.sha256(_limpar_texto(texto).encode("utf-8")).hexdigest()


class AuditorDOCX:
    """Inventário estrutural e integridade de conteúdo do arquivo DOCX."""

    def __init__(self, dados_docx: bytes):
        self.dados = dados_docx
        self.doc = Document(io.BytesIO(dados_docx))

    def _hash_cabecalhos_rodapes(self) -> str:
        texto = ""
        for section in self.doc.sections:
            for elemento in (section.header, section.footer):
                texto += "".join(p.text for p in elemento.paragraphs)
                for tabela in elemento.tables:
                    for linha in tabela.rows:
                        for celula in linha.cells:
                            texto += "".join(p.text for p in celula.paragraphs)
        return _sha(texto)

    def _hash_imagens(self) -> dict[str, str]:
        with zipfile.ZipFile(io.BytesIO(self.dados), "r") as arquivo:
            return {
                info.filename: hashlib.sha256(arquivo.read(info)).hexdigest()
                for info in arquivo.infolist()
                if info.filename.startswith("word/media/") and not info.is_dir()
            }

    def inventariar(self) -> dict:
        texto_p = "".join(p.text for p in self.doc.paragraphs)
        texto_t = "".join(
            p.text for tabela in self.doc.tables
            for linha in tabela.rows for celula in linha.cells
            for p in celula.paragraphs
        )
        return {
            "hash_paragrafos": _sha(texto_p),
            "hash_tabelas": _sha(texto_t),
            "hash_header_footer": self._hash_cabecalhos_rodapes(),
            "hash_imagens": self._hash_imagens(),
            "paragrafos_textos": [p.text for p in self.doc.paragraphs],
            "contagens": {
                "paragrafos": len(self.doc.paragraphs),
                "tabelas": len(self.doc.tables),
                "secoes": len(self.doc.sections),
            },
            "estilos": dict(Counter(p.style.name for p in self.doc.paragraphs)),
        }


class ClassificadorDeTitulos:
    @staticmethod
    def avaliar(doc) -> dict:
        acoes = {"h1": 0, "h2": 0, "h3": 0, "captions": 0, "candidatos_manuais": []}
        regex_h1 = re.compile(r"^(\d+)\s+([A-ZÀ-Úa-zà-ú ]+)$")
        regex_h2 = re.compile(r"^(\d+\.\d+)\s+([A-ZÀ-Úa-zà-ú ]+)$")
        regex_h3 = re.compile(r"^(\d+\.\d+\.\d+)\s+([A-ZÀ-Úa-zà-ú ]+)$")
        regex_caption = re.compile(r"^(Tabela|Figura|Quadro)\s+\d+", re.IGNORECASE)

        for paragrafo in doc.paragraphs:
            texto = paragrafo.text.strip()
            if not texto:
                continue
            if regex_h3.match(texto) and len(texto) < 150:
                paragrafo.style = "Heading 3"
                acoes["h3"] += 1
                continue
            if regex_h2.match(texto) and len(texto) < 150:
                paragrafo.style = "Heading 2"
                acoes["h2"] += 1
                continue
            if regex_h1.match(texto) and len(texto) < 100:
                paragrafo.style = "Heading 1"
                acoes["h1"] += 1
                continue
            if regex_caption.match(texto) and len(texto) < 200:
                paragrafo.style = "Caption"
                acoes["captions"] += 1
                continue
            if texto.isupper() and len(texto.split()) <= 6 and not texto.endswith("."):
                if not paragrafo.style.name.startswith("Heading"):
                    acoes["candidatos_manuais"].append(texto)
        return acoes


class GeradorDeSumario:
    @staticmethod
    def injetar(doc) -> bool:
        """Insere campo TOC na abertura, preservando os parágrafos originais.

        O rótulo recebe TOC Heading, não Heading 1, para não virar item do sumário.
        A atualização do campo deve ser realizada pelo usuário no Word.
        """
        if not doc.paragraphs:
            return False
        primeiro = doc.paragraphs[0]
        titulo = primeiro.insert_paragraph_before("SUMÁRIO")
        titulo.style = "TOC Heading" if "TOC Heading" in doc.styles else "Normal"
        campo = primeiro.insert_paragraph_before()
        run = campo.add_run()
        inicio = OxmlElement("w:fldChar")
        inicio.set(qn("w:fldCharType"), "begin")
        instrucao = OxmlElement("w:instrText")
        instrucao.set(qn("xml:space"), "preserve")
        instrucao.text = 'TOC \\o "1-3" \\h \\z \\u'
        separador = OxmlElement("w:fldChar")
        separador.set(qn("w:fldCharType"), "separate")
        fim = OxmlElement("w:fldChar")
        fim.set(qn("w:fldCharType"), "end")
        for elemento in (inicio, instrucao, separador, fim):
            run._r.append(elemento)
        return True


class FormatadorDOCX:
    @staticmethod
    def processar(dados_docx: bytes, perfil_nome: str, *, modo_conservador: bool = True,
                 gerar_sumario: bool = False) -> tuple[bytes, dict, dict]:
        if perfil_nome not in PERFIS:
            raise ValueError("Perfil de formatação inválido.")
        doc = Document(io.BytesIO(dados_docx))
        perfil = PERFIS[perfil_nome]
        log_formatacao = {"estilos_preservados": set(), "estilos_desconhecidos_ignorados": set()}

        for secao in doc.sections:
            secao.page_width = Cm(perfil["page"]["width"])
            secao.page_height = Cm(perfil["page"]["height"])
            secao.top_margin = Cm(perfil["page"]["top"])
            secao.bottom_margin = Cm(perfil["page"]["bottom"])
            secao.left_margin = Cm(perfil["page"]["left"])
            secao.right_margin = Cm(perfil["page"]["right"])

        for nome, regras in perfil["styles"].items():
            try:
                estilo = doc.styles[nome]
            except KeyError:
                continue
            estilo.font.name = regras["font"]
            estilo.font.size = Pt(regras["size"])
            # Ativa parâmetros já definidos no perfil original, inclusive
            # alinhamento, negrito e maiúsculas, antes não aplicados no Colab.
            estilo.font.bold = regras["bold"]
            if "all_caps" in regras:
                estilo.font.all_caps = regras["all_caps"]
            formato = estilo.paragraph_format
            formato.line_spacing = regras["spacing"]
            formato.first_line_indent = Cm(regras["indent"])
            formato.alignment = regras["align"]

        acoes = ClassificadorDeTitulos.avaliar(doc)
        acoes["toc_injetado"] = GeradorDeSumario.injetar(doc) if gerar_sumario else False

        for paragrafo in doc.paragraphs:
            nome = paragrafo.style.name
            if nome in ESTILOS_PROTEGIDOS:
                log_formatacao["estilos_preservados"].add(nome)
                continue
            if nome == "Normal":
                continue
            if modo_conservador:
                log_formatacao["estilos_desconhecidos_ignorados"].add(nome)
            else:
                paragrafo.style = "Normal"

        destino = io.BytesIO()
        doc.save(destino)
        return destino.getvalue(), acoes, log_formatacao


class RelatorioConformidade:
    @staticmethod
    def gerar_docx(nome_arquivo: str, perfil_nome: str, anterior: dict,
                   posterior: dict, acoes: dict, log_fmt: dict) -> bytes:
        certificado = Document()
        certificado.add_heading("Certificado de Auditoria e Conformidade (V1.1)", 0)
        certificado.add_paragraph(f"Arquivo: {nome_arquivo}")

        # Comparação específica para não classificar o sumário opcional
        # (um título e um campo novo) como perda de texto do documento.
        texto_pos = posterior["paragrafos_textos"]
        if acoes["toc_injetado"]:
            texto_pos = texto_pos[2:]
        conteudo_preservado = anterior["paragrafos_textos"] == texto_pos
        estado_paragrafos = "Preservado" if conteudo_preservado else "Divergente — revisar"

        certificado.add_heading("1. Integridade de Dados", level=1)
        tabela = certificado.add_table(rows=1, cols=2)
        tabela.style = "Table Grid"
        tabela.rows[0].cells[0].text = "Métrica (SHA256 / conteúdo)"
        tabela.rows[0].cells[1].text = "Status de preservação"
        linhas = [
            ("Texto dos parágrafos", estado_paragrafos),
            ("Texto das tabelas", "Preservado" if anterior["hash_tabelas"] == posterior["hash_tabelas"] else "Divergente — revisar"),
            ("Cabeçalhos e rodapés", "Preservado" if anterior["hash_header_footer"] == posterior["hash_header_footer"] else "Divergente — revisar"),
            ("Mídias / imagens", f"Preservadas ({len(anterior['hash_imagens'])} arquivos)" if anterior["hash_imagens"] == posterior["hash_imagens"] else "Divergente — revisar"),
        ]
        for titulo, resultado in linhas:
            linha = tabela.add_row()
            linha.cells[0].text = titulo
            linha.cells[1].text = resultado

        certificado.add_heading("2. Formatações Aplicadas", level=1)
        for descricao in (
            f"Perfil selecionado: {perfil_nome}",
            f"Títulos nível 1 classificados: {acoes['h1']}",
            f"Títulos nível 2 classificados: {acoes['h2']}",
            f"Títulos nível 3 classificados: {acoes['h3']}",
            f"Legendas classificadas: {acoes['captions']}",
        ):
            certificado.add_paragraph(descricao, style="List Bullet")

        certificado.add_heading("3. Estilos Protegidos e Preservados", level=1)
        encontrados = log_fmt["estilos_preservados"] | log_fmt["estilos_desconhecidos_ignorados"]
        if encontrados:
            for estilo in sorted(encontrados):
                certificado.add_paragraph(estilo, style="List Bullet")
        else:
            certificado.add_paragraph("Nenhum estilo especial identificado além do estilo Normal.")

        certificado.add_heading("4. Alertas para Revisão Manual", level=1)
        if acoes.get("toc_injetado"):
            certificado.add_paragraph(
                "SUMÁRIO INSERIDO: no Microsoft Word, clique com o botão direito sobre o sumário e "
                "escolha 'Atualizar Campo' > 'Atualizar o índice inteiro'. O sumário depende "
                "dessa atualização para preencher os itens e páginas.",
                style="List Bullet",
            )
        if acoes["candidatos_manuais"]:
            certificado.add_paragraph("Possíveis títulos não numerados identificados para revisão:")
            for candidato in acoes["candidatos_manuais"]:
                certificado.add_paragraph(f'"{candidato}"', style="List Bullet")
        else:
            certificado.add_paragraph("Nenhum candidato a título manual detectado.")
        certificado.add_paragraph(
            "A auditoria compara texto e arquivos de imagem. Ela não certifica conformidade "
            "completa à ABNT nem garante a preservação de todos os elementos complexos do Word."
        )

        destino = io.BytesIO()
        certificado.save(destino)
        return destino.getvalue()


@dataclass(frozen=True)
class ResultadoFormatacao:
    formatado: bytes
    certificado: bytes
    acoes: dict
    integridade: dict


def formatar_relatorio(dados_docx: bytes, nome_arquivo: str, perfil: str,
                       gerar_sumario: bool = False) -> ResultadoFormatacao:
    """Processa um único arquivo inteiramente em memória, sem gravar histórico."""
    if not dados_docx:
        raise ValueError("Selecione um arquivo Word válido.")
    if len(dados_docx) > MAX_BYTES:
        raise ValueError("O documento excede o limite de 25 MB.")
    if not nome_arquivo.lower().endswith(".docx"):
        raise ValueError("A ferramenta aceita somente documentos .docx.")
    if not zipfile.is_zipfile(io.BytesIO(dados_docx)):
        raise ValueError("O arquivo não contém uma estrutura DOCX válida.")

    anterior = AuditorDOCX(dados_docx).inventariar()
    arquivo_formatado, acoes, log_fmt = FormatadorDOCX.processar(
        dados_docx, perfil, modo_conservador=True, gerar_sumario=gerar_sumario
    )
    posterior = AuditorDOCX(arquivo_formatado).inventariar()
    certificado = RelatorioConformidade.gerar_docx(
        nome_arquivo, perfil, anterior, posterior, acoes, log_fmt
    )
    integridade = {
        "paragrafos": anterior["paragrafos_textos"] == (
            posterior["paragrafos_textos"][2:] if acoes["toc_injetado"]
            else posterior["paragrafos_textos"]
        ),
        "tabelas": anterior["hash_tabelas"] == posterior["hash_tabelas"],
        "cabecalhos_rodapes": anterior["hash_header_footer"] == posterior["hash_header_footer"],
        "imagens": anterior["hash_imagens"] == posterior["hash_imagens"],
    }
    return ResultadoFormatacao(arquivo_formatado, certificado, acoes, integridade)


def render_formatador_relatorios():
    """Formulário chamado pelo st.dialog da página do Módulo 6."""
    import streamlit as st

    st.caption("Envie um documento .docx, selecione o padrão e baixe o arquivo formatado e o certificado de auditoria.")
    upload = st.file_uploader(
        "Documento Word (.docx)",
        type=["docx"],
        key="wordtools_upload_relatorio",
        help="Somente .docx. Máximo: 25 MB. O arquivo original não será modificado.",
    )
    perfil = st.selectbox(
        "Perfil de formatação",
        options=list(PERFIS),
        format_func=lambda chave: ROTULOS_PERFIS[chave],
        key="wordtools_perfil_relatorio",
    )
    gerar_sumario = st.checkbox(
        "Inserir sumário automático no início do documento",
        value=False,
        key="wordtools_sumario_relatorio",
    )
    if gerar_sumario:
        st.caption("O campo do sumário será inserido. Atualize-o no Microsoft Word para exibir os títulos e números de página.")
    st.caption("Modo conservador ativo: estilos especiais são preservados; os estilos básicos seguem o perfil escolhido.")

    if upload is None:
        st.session_state.pop("wordtools_resultado_relatorio", None)
        st.info("Selecione um documento para habilitar a formatação.")
        return

    dados = upload.getvalue()
    assinatura = hashlib.sha256(dados + perfil.encode("utf-8") + bytes([int(gerar_sumario)])).hexdigest()

    if st.button(
        "📄 Formatar Relatório",
        type="primary",
        use_container_width=True,
        key="wordtools_executar_relatorio",
    ):
        try:
            with st.spinner("Formatando documento e verificando integridade..."):
                resultado = formatar_relatorio(dados, upload.name, perfil, gerar_sumario)
            st.session_state["wordtools_resultado_relatorio"] = (assinatura, resultado, upload.name)
        except (ValueError, zipfile.BadZipFile) as erro:
            st.session_state.pop("wordtools_resultado_relatorio", None)
            st.error(str(erro))
        except Exception:
            st.session_state.pop("wordtools_resultado_relatorio", None)
            st.error("Não foi possível processar este documento. Verifique se o arquivo .docx está íntegro e tente novamente.")

    item = st.session_state.get("wordtools_resultado_relatorio")
    if not item or item[0] != assinatura:
        return

    _assinatura, resultado, nome_original = item
    st.success("Documento processado. Os dois arquivos estão disponíveis para download.")
    if not all(resultado.integridade.values()):
        st.warning("A auditoria encontrou divergências de conteúdo. Confira o certificado antes de utilizar o documento.")
    st.caption(
        f"Classificações: {resultado.acoes['h1']} título(s) nível 1, "
        f"{resultado.acoes['h2']} nível 2, "
        f"{resultado.acoes['h3']} nível 3 e "
        f"{resultado.acoes['captions']} legenda(s)."
    )

    nome_limpo = re.sub(r"[^a-zA-Z0-9._-]+", "_", nome_original)
    col1, col2 = st.columns(2)
    with col1:
        st.download_button(
            "⬇️ Baixar relatório formatado",
            data=resultado.formatado,
            file_name=f"FORMATADO_{nome_limpo}",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
            key="wordtools_download_formatado",
        )
    with col2:
        st.download_button(
            "⬇️ Baixar certificado",
            data=resultado.certificado,
            file_name="Certificado_Conformidade.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
            key="wordtools_download_certificado",
        )
