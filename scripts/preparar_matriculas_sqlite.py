"""Converte todas as abas válidas de matriculas.xls/.xlsx para SQLite indexado.

Execução na raiz do repositório:
    python scripts/preparar_matriculas_sqlite.py
ou:
    python scripts/preparar_matriculas_sqlite.py dados/matriculas.xls dados/matriculas.sqlite

O banco contém dados pessoais dos clientes e NÃO deve ir a repositórios públicos.
"""
from __future__ import annotations

import argparse
import hashlib
import os
from pathlib import Path
import re
import sqlite3
import unicodedata
import uuid

CAMPOS = ("CATEGORIA", "NUM_LIGACAO", "BAIRRO", "COD_LATITUDE", "COD_LONGITUDE", "END_LIGACAO")


def normalizar(campo):
    texto = unicodedata.normalize("NFKD", str(campo if campo is not None else "").upper())
    return re.sub(r"[^A-Z0-9]", "", "".join(c for c in texto if not unicodedata.combining(c)))


def texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def linhas_por_aba(caminho):
    with caminho.open("rb") as arq:
        assinatura = arq.read(8)
    if assinatura.startswith(b"PK"):
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise RuntimeError("Instale openpyxl: pip install openpyxl") from e
        entrada = caminho.open("rb")
        try:
            # Aceita arquivo XLSX com extensão .xls incorreta.
            livro = load_workbook(entrada, read_only=True, data_only=True)
            for aba in livro.worksheets:
                linhas = aba.iter_rows(values_only=True)
                cabecalho = next(linhas, None)
                if cabecalho is None:
                    continue
                m = {normalizar(c): i for i, c in enumerate(cabecalho)}
                if not all(normalizar(c) in m for c in CAMPOS):
                    continue
                indices = [m[normalizar(c)] for c in CAMPOS]
                yield aba.title, (tuple(linha[i] if i < len(linha) else "" for i in indices) for linha in linhas)
        finally:
            if "livro" in locals():
                livro.close()
            entrada.close()
    elif assinatura.startswith(b"\xd0\xcf\x11\xe0"):
        try:
            import xlrd
        except ImportError as e:
            raise RuntimeError("Instale xlrd: pip install 'xlrd>=2.0.1'") from e
        livro = xlrd.open_workbook(str(caminho), on_demand=True)
        try:
            for aba in livro.sheets():
                if not aba.nrows:
                    continue
                m = {normalizar(c): i for i, c in enumerate(aba.row_values(0))}
                if not all(normalizar(c) in m for c in CAMPOS):
                    continue
                indices = [m[normalizar(c)] for c in CAMPOS]
                def iterar(a=aba, cols=indices):
                    for idx in range(1, a.nrows):
                        yield tuple(a.cell_value(idx, c) if c < a.ncols else "" for c in cols)
                yield aba.name, iterar()
        finally:
            livro.release_resources()
    else:
        raise ValueError("Formato Excel inválido. O arquivo deve ser XLS ou XLSX.")


def converter(origem: Path, destino: Path):
    if not origem.is_file():
        raise FileNotFoundError(f"Arquivo de origem não encontrado: {origem}")
    if origem.resolve() == destino.resolve():
        raise ValueError("Origem e destino não podem ser o mesmo arquivo.")
    hash_excel = hashlib.sha256()
    with origem.open("rb") as arq:
        for bloco in iter(lambda: arq.read(2 * 1024 * 1024), b""):
            hash_excel.update(bloco)
    destino.parent.mkdir(parents=True, exist_ok=True)
    temporario = destino.with_name(destino.name + "." + uuid.uuid4().hex + ".tmp")
    total, abas_validas = 0, 0
    try:
        os.close(os.open(temporario, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600))
        with sqlite3.connect(temporario) as con:
            con.execute("PRAGMA journal_mode=OFF")
            con.execute("PRAGMA synchronous=OFF")
            con.execute("CREATE TABLE matriculas (CATEGORIA TEXT, NUM_LIGACAO TEXT NOT NULL, "
                        "BAIRRO TEXT, COD_LATITUDE TEXT, COD_LONGITUDE TEXT, END_LIGACAO TEXT)")
            con.execute("CREATE TABLE metadados (chave TEXT PRIMARY KEY, valor TEXT)")
            for nome, linhas in linhas_por_aba(origem):
                abas_validas += 1
                lote = []
                qtd_aba = 0
                for linha in linhas:
                    celulas = [texto(v) for v in linha]
                    codigo = re.sub(r"\.0+$", "", celulas[1]) if re.fullmatch(r"\d+\.0+", celulas[1]) else celulas[1]
                    celulas[1] = codigo
                    if not codigo:
                        continue
                    lote.append(celulas)
                    if len(lote) >= 5000:
                        con.executemany("INSERT INTO matriculas VALUES (?,?,?,?,?,?)", lote)
                        total += len(lote)
                        qtd_aba += len(lote)
                        lote.clear()
                if lote:
                    con.executemany("INSERT INTO matriculas VALUES (?,?,?,?,?,?)", lote)
                    total += len(lote)
                    qtd_aba += len(lote)
                print(f"Aba {nome}: {qtd_aba:,} matrículas processadas".replace(",", "."))
            if not abas_validas:
                raise ValueError("Não foram encontradas abas com as seis colunas necessárias.")
            con.execute("CREATE INDEX idx_num_ligacao ON matriculas(NUM_LIGACAO)")
            con.executemany("INSERT INTO metadados(chave, valor) VALUES (?, ?)", [
                ("versao", "3"), ("fonte", origem.name), ("sha256", hash_excel.hexdigest()),
                ("quantidade", str(total)), ("abas_validas", str(abas_validas)),
            ])
            con.commit()
        os.replace(temporario, destino)
        os.chmod(destino, 0o600)
        print(f"\nConcluído: {total:,} registros em {abas_validas} aba(s).".replace(",", "."))
        print(f"SQLite: {destino} ({destino.stat().st_size / 1024**2:.1f} MiB)")
        print("AVISO: O banco contém informações cadastrais. Armazene apenas em ambiente privado e autorizado.")
    finally:
        if temporario.exists():
            temporario.unlink()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Pré-indexar matrículas do COI para consultas rápidas")
    parser.add_argument("origem", nargs="?", default="dados/matriculas.xls")
    parser.add_argument("destino", nargs="?", default="dados/matriculas.sqlite")
    args = parser.parse_args()
    converter(Path(args.origem), Path(args.destino))
